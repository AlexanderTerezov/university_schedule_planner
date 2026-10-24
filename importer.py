"""Step 1 - read the two Excel sheets and turn them into clean courses and sessions.

Nothing is silently invented or dropped: every skipped, kept-with-a-caveat or
auto-corrected record is logged as an Issue so somebody can fix the source file.
"""
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from functools import partial

import openpyxl

import config

LECTURE, EXERCISE = "Лекция", "Упражнение"
DAYS = {"понеделник": 0, "вторник": 1, "сряда": 2, "четвъртък": 3,
        "петък": 4, "събота": 5, "неделя": 6}

# Latin letters that look like Cyrillic ones. Typed by mistake they create "different"
# rooms/specialties (146A vs 146А), which would hide real clashes.
_LOOKALIKE = str.maketrans("ABCEHKMOPTX", "АВСЕНКМОРТХ")


@dataclass
class Issue:
    level: str    # ERROR = record skipped | WARNING = kept, please check | INFO = corrected / nothing to do
    sheet: str
    row: int
    subject: str
    message: str


@dataclass(frozen=True)
class Requirement:
    """One preference from the "Изисквания" column (all of them are soft).

    at / not_at   -> the class starts / does not start at this time
    after / before -> the class starts at or after / ends at or before this time
    on_days       -> the class is on one of these days (0 = Monday)
    not_day       -> the class is not on this day
    Times are minutes since midnight.
    """
    kind: str
    value: object
    raw: str

    def ok(self, weekday, start, end):
        """Is a class on this weekday (0 = Monday), from `start` to `end` (minutes), within the preference?"""
        if self.kind == "at":
            return start == self.value
        if self.kind == "not_at":
            return start != self.value
        if self.kind == "after":
            return start >= self.value
        if self.kind == "before":
            return end <= self.value
        if self.kind == "on_days":
            return weekday in self.value
        return weekday != self.value          # not_day


@dataclass
class Course:
    kind: str            # LECTURE or EXERCISE
    sheet: str
    row: int
    subject: str
    hours: int           # total hours for ONE group
    cohorts: tuple       # e.g. ("СДТНГ 3", "ГГПТТ 3") - all of them attend together
    groups: int          # each group gets the full number of hours
    duration: int        # "Продължение": length of one session in hours
    teachers: tuple      # all of them are needed for every session - unless each group has its own
    rooms: tuple         # in order of preference
    requirements: tuple

    @property
    def parallel(self):
        """As many teachers as groups: group i is taught by teacher i and all groups run at the same time."""
        return self.groups > 1 and len(self.teachers) == self.groups


@dataclass
class Session:
    id: int
    course: Course
    group: int           # 1..course.groups
    part: int            # 0, 1, 2... position among the sessions of this group
    length: int          # hours
    teachers: tuple      # who has to be present
    resized: bool = False  # True when a repair round cut it with a length other than the usual one


# ---------------------------------------------------------------------------
# Reading the workbook
# ---------------------------------------------------------------------------
def _fill_merged_cells(ws):
    """Give every cell of a merged block the block's value (data rows only)."""
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row < 3:          # the two header rows
            continue
        value = ws.cell(rng.min_row, rng.min_col).value
        ws.unmerge_cells(str(rng))
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                ws.cell(r, c).value = value


def _data_rows(ws, ncols):
    """Yield (excel_row, values) for non-empty rows below the 2-row header.

    Stops after 50 empty rows: Sheet2 claims to have a million rows."""
    blank = 0
    for r, values in enumerate(ws.iter_rows(min_row=3, max_col=ncols, values_only=True), start=3):
        if all(v is None for v in values):
            blank += 1
            if blank >= 50:
                return
        else:
            blank = 0
            yield r, values


def read_raw(path):
    """Return one dict per row of both sheets, with the two sheets' layouts unified."""
    wb = openpyxl.load_workbook(path)
    raw = []

    ws = wb[config.LECTURES_SHEET]
    _fill_merged_cells(ws)
    for r, v in _data_rows(ws, 10):
        raw.append(dict(kind=LECTURE, sheet=ws.title, row=r, subject=v[1], hours=[v[2]],
                        spec=v[3], groups=1, rooms=v[4:7], teacher=v[7], cont=v[8], req=v[9]))

    ws = wb[config.EXERCISES_SHEET]
    _fill_merged_cells(ws)
    for r, v in _data_rows(ws, 13):
        # Columns: №, Дисциплина, СУ, ПУ, ЛУ, Специалност и курс, Брой групи, Зала x3,
        # Преподавател, Продължение, Изисквания. hours = СУ + ПУ + ЛУ (there is no SUM column).
        raw.append(dict(kind=EXERCISE, sheet=ws.title, row=r, subject=v[1], hours=list(v[2:5]),
                        spec=v[5], groups=v[6], rooms=v[7:10], teacher=v[10], cont=v[11], req=v[12]))
    return raw


# ---------------------------------------------------------------------------
# Parsing single cells
# ---------------------------------------------------------------------------
def clean(value):
    return re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""


def _as_int(value):
    """int for numbers, None for an empty cell, ValueError for anything else (e.g. text)."""
    text = clean(value)
    if text == "":
        return None
    if isinstance(value, (int, float)) and value == int(value):
        return int(value)
    if text.isdigit():
        return int(text)
    raise ValueError(text)


def parse_cohorts(text, spelling, note):
    """'СДТНГ 3, ГГПТТ 3' -> ('СДТНГ 3', 'ГГПТТ 3'). Copes with ';', '.', 'АИУТ4' and known typos."""
    cohorts = []
    for name, year in re.findall(r"([А-Яа-яA-Za-z]+)\s*(\d)", text):
        key = name.upper().translate(_LOOKALIKE)
        fixed = config.SPECIALTY_ALIASES.get(key, key)
        if fixed != key:
            note("INFO", f"specialty {name} {year} read as {fixed} {year}")
        label = f"{spelling.setdefault(fixed, name if fixed == key else fixed)} {year}"
        if label not in cohorts:
            cohorts.append(label)
    return tuple(cohorts)


def parse_rooms(cells, note):
    """Three preference cells -> room names, best first. 'кабинет 312', '146 A', '320/321' are handled.
    Text without a number ('в мемф', 'кабинет') is kept as a room name too."""
    rooms = []
    for cell in cells:
        text = clean(cell)
        if "+" in text:
            note("WARNING", f"room cell {text!r} uses '+' (several rooms at once?) - treated as alternatives")
        for part in re.split(r"[,/+]", text):
            if not part.strip():
                continue
            room = re.sub(r"(?i)кабинет", "", part).replace(" ", "").upper().translate(_LOOKALIKE)
            if not any(ch.isdigit() for ch in room):
                # no room number ("в мемф", "кабинет"): any text is accepted as a room name of its own
                room = re.sub(r"\s+", " ", part).strip().upper().translate(_LOOKALIKE)
            room = config.ROOM_ALIASES.get(room, room)
            if room not in rooms:
                rooms.append(room)
    return tuple(rooms)


def parse_requirements(cell, note):
    """'8, ! 9:00, > 10, < 16:30, вторник, ! сряда' -> Requirement objects.

    Each comma- or semicolon-separated token is one preference:
      8  /  8:00        the class starts at 8:00
      ! 8  /  ! 8:00    the class does not start at 8:00
      > 10  /  > 10:30  the class starts at or after that time
      < 16  /  < 16:30  the class ends at or before that time
      вторник           the class is on that day
      ! сряда           the class is not on that day

    Times can be written as H, HH, H:MM or HH:MM (also H.MM, and HH:MM:SS as Excel
    sometimes stores a bare time). "!" cannot be combined with < or >.
    Several positive days ('вторник, сряда') become ONE preference: any of them is fine.

    Warnings are logged for invalid times (e.g. 25:70) and for exact start times that
    can never be met because classes only start on multiples of SLOT_MINUTES."""
    reqs, days, day_raw = [], set(), []
    for token in re.split(r"[,;]", clean(cell)):
        raw = token.strip()
        if not raw:
            continue
        negated = raw.startswith("!")
        body = raw.lstrip("!").strip().lower()
        op = ""
        if body[:1] and body[:1] in "<>":
            op, body = body[0], body[1:].strip()
        time_match = re.fullmatch(r"(\d{1,2})(?:[:.](\d{2})(?::\d{2})?)?", body)

        if body in DAYS and not op:
            if negated:
                reqs.append(Requirement("not_day", DAYS[body], raw))
            else:
                days.add(DAYS[body])
                day_raw.append(raw)
        elif time_match and not (negated and op):
            hours, mins = int(time_match[1]), int(time_match[2] or 0)
            if hours > 24 or mins > 59 or (hours == 24 and mins):
                note("WARNING", f"not a valid time, ignored: {raw!r}")
                continue
            minutes = hours * 60 + mins
            kind = {"": "not_at" if negated else "at", ">": "after", "<": "before"}[op]
            if kind == "at" and (minutes - config.DAY_START_MIN) % config.SLOT_MINUTES:
                note("WARNING", f"{raw!r} can never be met: classes only start on "
                                f"multiples of {config.SLOT_MINUTES} min from {config.DAY_START}")
            reqs.append(Requirement(kind, minutes, raw))
        else:
            note("WARNING", f"requirement not understood, ignored: {raw!r}")
    if days:
        reqs.append(Requirement("on_days", frozenset(days), ", ".join(day_raw)))
    return tuple(reqs)


# --- Teachers: the same person is written many ways ("доц. Е. Крайчев", "доц. д-р Евгени Крайчев")
_TITLE = re.compile(
    r"(?<![\w-])(?:проф|доц|гл|ас|ст\.?\s*пр|д-р|дн|дбн|инж|маг|др|докторант)\.?"
    r"(?=\s|\.|$|(?-i:[А-ЯA-Z]))", re.IGNORECASE)


def parse_teachers(text):
    """Cell -> list of (given_name, surname), lower case. Titles are dropped; the given name may be an initial."""
    pairs = []
    # split BEFORE collapsing whitespace: a line break or two or more spaces separate people
    for part in re.split(r"[,;/&\n]|\s{2,}", str(text or "")):
        tokens = [t for t in re.split(r"[.\s]+", _TITLE.sub(" ", part)) if t and t not in "-–"]
        if tokens:
            pairs.append((tokens[0].lower() if len(tokens) > 1 else "", tokens[-1].lower()))
    return pairs


def build_teacher_index(pairs, issues):
    """Map every (given, surname) spelling to one display name per person.

    Same surname + given names where one starts with the other ("е" / "евгени") = same person.
    A missing given name matches any person with that surname."""
    people = defaultdict(list)           # surname -> given names of distinct people
    index = {}
    for given, surname in sorted(set(pairs), key=lambda p: -len(p[0])):   # full names first
        same = [g for g in people[surname] if g.startswith(given) or given.startswith(g)]
        if not same:
            people[surname].append(given)
            same = [given]
        elif len(same) > 1:
            issues.append(Issue("INFO", "teachers", 0, "", f"'{given} {surname}' could be several people; "
                                f"assumed the same as '{same[0]} {surname}'"))
        g = same[0]
        shown = f"{g.capitalize()}{'.' if len(g) <= 2 else ''} " if g else ""
        index[(given, surname)] = shown + surname.capitalize()
    return index


# ---------------------------------------------------------------------------
# Building courses
# ---------------------------------------------------------------------------
def _build_course(rec, subject, teacher_index, spelling, note):
    if not subject:
        note("ERROR", "no subject name - record skipped")
        return None
    cohorts = parse_cohorts(clean(rec["spec"]), spelling, note)
    if not cohorts:
        note("ERROR", f"no specialty/course found in {clean(rec['spec'])!r} - record skipped")
        return None

    try:
        hours = sum(_as_int(cell) or 0 for cell in rec["hours"])
        groups = _as_int(rec["groups"])
        duration = _as_int(rec["cont"])
    except ValueError as err:
        note("ERROR", f"text where a number is expected ({err}) - record skipped")
        return None

    if hours == 0:
        note("INFO", "0 hours - nothing to schedule")
        return None
    for name, value, default in (("number of groups", groups, config.DEFAULT_GROUPS),
                                 ("continuation length", duration, config.DEFAULT_CONTINUATION)):
        if value is None and default is None:
            note("ERROR", f"{name} is missing - record skipped")
            return None
        if value is None:
            note("INFO", f"{name} is missing - default {default} used")
    groups = groups if groups is not None else config.DEFAULT_GROUPS
    duration = duration if duration is not None else config.DEFAULT_CONTINUATION
    if groups < 1 or duration < 1:
        note("ERROR", "groups and continuation must be at least 1 - record skipped")
        return None

    teachers = tuple(dict.fromkeys(teacher_index[p] for p in parse_teachers(rec["teacher"])))
    if not teachers:
        note("WARNING", "no teacher - teacher clashes cannot be checked for this record")
    elif groups > 1 and len(teachers) > 1:
        if len(teachers) == groups:
            note("INFO", f"{groups} groups and {groups} teachers - each group has its own teacher, groups run in parallel")
        else:
            note("WARNING", f"{len(teachers)} teachers and {groups} groups - assumed ALL teachers attend EVERY session")
    rooms = parse_rooms(rec["rooms"], note)
    if not rooms:
        note("WARNING", "no room given")
    if rooms and groups > 1 and len(teachers) == groups and len(rooms) < groups:
        note("WARNING", f"{groups} groups run in parallel but only {len(rooms)} room(s) listed - "
                        f"they cannot all be placed")

    return Course(rec["kind"], rec["sheet"], rec["row"], subject, hours, cohorts, groups, duration,
                  teachers, rooms, parse_requirements(rec["req"], note))


def _log(issues, rec, subject, level, message):
    issues.append(Issue(level, rec["sheet"], rec["row"], subject, message))


def load_courses(path):
    """Read the workbook. Returns (courses, issues, rows_read) where rows_read counts records per kind."""
    raw = read_raw(path)
    issues = []
    all_pairs = [p for rec in raw for p in parse_teachers(rec["teacher"])]
    teacher_index = build_teacher_index(all_pairs, issues)

    courses, spelling, rows_read = [], {}, Counter()
    for rec in raw:
        subject = clean(rec["subject"])
        if not (subject or clean(rec["spec"]) or clean(rec["teacher"])):
            continue                                     # the totals row at the bottom of Sheet2
        rows_read[rec["kind"]] += 1
        note = partial(_log, issues, rec, subject)
        course = _build_course(rec, subject, teacher_index, spelling, note)
        if course:
            courses.append(course)

    seen = Counter(c for course in courses for c in course.cohorts)
    for cohort, n in seen.items():
        if n <= 2:
            issues.append(Issue("INFO", "cohorts", 0, "", f"'{cohort}' appears in only {n} course(s) - typo?"))
    return courses, issues, rows_read


def _subject_key(name):
    """Subject name reduced to letters and digits, so spacing, case and punctuation do not matter."""
    return re.sub(r"[\W_]+", "", name.lower().translate(str.maketrans("abcehkmoptx", "авсенкмортх")))


def link_lectures(courses):
    """Pairs (lecture, exercise) of the same subject for the same students.

    Matched by subject name AND at least one shared cohort, so a subject with the same
    name in an unrelated specialty is not linked by accident."""
    lectures = defaultdict(list)
    for course in courses:
        if course.kind == LECTURE:
            lectures[_subject_key(course.subject)].append(course)
    return [(lecture, course)
            for course in courses if course.kind == EXERCISE
            for lecture in lectures.get(_subject_key(course.subject), ())
            if set(lecture.cohorts) & set(course.cohorts)]


def expand_sessions(courses):
    """Split every course into sessions: 14 h with continuation 3 -> 3+3+3+3+2.
    Every group of a course gets the full number of hours."""
    sessions = []
    for course in courses:
        for group in range(1, course.groups + 1):
            teachers = (course.teachers[group - 1],) if course.parallel else course.teachers
            left, part = course.hours, 0
            while left > 0:
                length = min(course.duration, left)
                sessions.append(Session(len(sessions), course, group, part, length, teachers))
                left -= length
                part += 1
    return sessions