"""Render the solver output (timetable.xlsx) as readable timetables, one sheet per cohort.

Usage:
    python render_timetable.py                              -> timetable_view.xlsx, one sheet per cohort
    python render_timetable.py timetable.xlsx view.xlsx
    python render_timetable.py timetable.xlsx view.xlsx --by teacher     (one sheet per teacher)
    python render_timetable.py timetable.xlsx view.xlsx --by room        (one sheet per room)

Each sheet is a grid: one row per date (grouped in weeks), one column per hour.
A class is a coloured block spanning its hours, with subject, teacher, type and room.
Two classes at the same time (e.g. two exercise groups) are stacked in the same day.
Below the grid, classes of that cohort that the solver could not place are listed.
It only reads the Excel output, so it does not need the solver or OR-Tools.
"""
import argparse
import re
from collections import defaultdict
from datetime import timedelta
from math import gcd
from typing import Any, cast

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font as _Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

try:                                    # use the project's day and lunch settings when config.py is next to this file
    import config
    DAY_START, DAY_END = config.DAY_START_MIN, config.DAY_END_MIN
    LUNCH_START, LUNCH_MINUTES = config.LUNCH_START_MIN, config.LUNCH_MINUTES
except Exception:                       # pragma: no cover - standalone fallback
    DAY_START, DAY_END, LUNCH_START, LUNCH_MINUTES = 8 * 60, 20 * 60, 12 * 60, 60

def Font(**kw):
    """Every cell uses Calibri explicitly, so the look does not depend on the viewer's default font."""
    return _Font(name="Calibri", **kw)


# --- Text shown in the workbook (change here to translate) ---------------------------------
L = {
    "title": "Разпис на занятията",
    "period": "Период",
    "week": "Седм.",
    "day": "Ден",
    "lecture": "Лекция",
    "exercise": "Упражнение",
    "lunch": "Обедна почивка",
    "warn": "⚠ нарушено изискване (виж бележката)",
    "unplaced": "Неразпределени занятия",
    "subject": "Предмет", "type": "Вид", "group": "Група", "parts": "Части", "hours": "Часове",
    "teacher": "Преподавател", "cohorts": "Специалности",
    "no_room": "без зала",
    "none_room": "Без зала",
    "overview": "Преглед",
    "placed_n": "Разпределени занятия", "placed_h": "Разпределени часове",
    "unplaced_n": "Неразпределени занятия", "unplaced_h": "Неразпределени часове",
    "placed_pct": "% разпределени часове",
    "by": {"cohort": "Специалност и курс", "teacher": "Преподавател", "room": "Зала"},
}
DAY_NAMES = ["Понеделник", "Вторник", "Сряда", "Четвъртък", "Петък", "Събота", "Неделя"]
TYPE_SHORT = {"Лекция": "Л", "Упражнение": "У"}

# --- Look ------------------------------------------------------------------------------------
FILL = {
    "Лекция": PatternFill("solid", fgColor="DCE9F7"),
    "Упражнение": PatternFill("solid", fgColor="DDF0DD"),
}
EDGE = {"Лекция": "5B8FC7", "Упражнение": "5FA55F"}
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
LUNCH_FILL = PatternFill("solid", fgColor="E7E7E7")
WEEKEND_FILL = PatternFill("solid", fgColor="F6F6F6")
DAY_FILL = PatternFill("solid", fgColor="EEF2F7")
THIN = Side(style="thin", color="D0D0D0")
MEDIUM = Side(style="medium", color="1F3A5F")
FIRST_COL = 3                           # A = week, B = day, hours start in C
COL_WIDTH_HOUR = 16
ROW_HEIGHT_CLASS = 54
ROW_HEIGHT_EMPTY = 20
FONT_SIZE = 8


# ---------------------------------------------------------------------------------------------
# Reading the solver output
# ---------------------------------------------------------------------------------------------
def minutes(text):
    hours, mins = str(text).split(":")[:2]
    return int(hours) * 60 + int(mins)


def as_int(value: Any) -> int:
    return int(value)


def as_date(value: Any):
    return pd.Timestamp(value).date()


def split_list(value):
    return [part.strip() for part in str(value).split(",") if part.strip()] if pd.notna(value) else []


def read_sessions(path):
    xl = pd.ExcelFile(path)
    placed = pd.read_excel(xl, "Timetable")
    try:
        unplaced = pd.read_excel(xl, "Unscheduled")
    except ValueError:
        unplaced = pd.DataFrame()

    # exercises of a subject that has several groups show the group number; single-group ones do not
    multi_group = set()
    for df in (placed, unplaced):
        if not df.empty:
            for (kind, subject, cohorts), grp in df.groupby(["Type", "Subject", "Cohorts"]):
                if grp["Group"].max() > 1:
                    multi_group.add((kind, subject, cohorts))

    sessions = []
    for r in placed.itertuples(index=False):
        sessions.append(dict(
            kind=r.Type, subject=r.Subject, cohorts=split_list(r.Cohorts), group=as_int(r.Group), part=as_int(r.Part),
            hours=as_int(r[5]), teachers=split_list(r.Teachers), day=as_date(r.Date),
            start=minutes(r.Start), end=minutes(r.End),
            room="" if pd.isna(r.Room) or r.Room == "NOT ASSIGNED" else str(r.Room),
            warnings=[] if pd.isna(r[12]) else [w.strip() for w in str(r[12]).split(";") if w.strip()],
            multi=(r.Type, r.Subject, r.Cohorts) in multi_group))

    missing = []
    for r in unplaced.itertuples(index=False):
        missing.append(dict(
            kind=r.Type, subject=r.Subject, cohorts=split_list(r.Cohorts), group=as_int(r.Group), part=as_int(r.Part),
            hours=as_int(r[5]), teachers=split_list(r.Teachers),
            multi=(r.Type, r.Subject, r.Cohorts) in multi_group))
    return sessions, missing


# ---------------------------------------------------------------------------------------------
# Choosing the grid
# ---------------------------------------------------------------------------------------------
def make_grid(sessions) -> dict[str, Any]:
    starts = [s["start"] for s in sessions] + [DAY_START]
    ends = [s["end"] for s in sessions] + [DAY_END]
    step = 60
    for value in starts + ends + [LUNCH_START, LUNCH_START + LUNCH_MINUTES]:
        step = gcd(step, value % 60 or 60)
    step = max(step, 5)
    t0 = (min(starts) // 60) * 60
    t1 = -(-max(ends) // 60) * 60
    days = [s["day"] for s in sessions]
    first = min(days) - timedelta(days=min(days).weekday())
    last = max(days) + timedelta(days=6 - max(days).weekday())
    weeks = [first + timedelta(days=7 * i) for i in range((last - first).days // 7 + 1)]
    return dict(step=step, t0=t0, t1=t1, ncols=(t1 - t0) // step, weeks=weeks,
                first=min(days), last=max(days))


def clock(m):
    return f"{m // 60:02d}:{m % 60:02d}"


def short_list(items, limit):
    items = list(items)
    return ", ".join(items) if len(items) <= limit else f"{', '.join(items[:limit - 1])} +{len(items) - limit + 1}"


def assign_lanes(items):
    """Stack classes that overlap in time. Returns ({id(item): lane}, number of lanes)."""
    lane_end, lane_of = [], {}
    for item in sorted(items, key=lambda s: (s["start"], s["end"])):
        for i, end in enumerate(lane_end):
            if end <= item["start"]:
                lane_of[id(item)] = i
                lane_end[i] = item["end"]
                break
        else:
            lane_of[id(item)] = len(lane_end)
            lane_end.append(item["end"])
    return lane_of, len(lane_end)


# ---------------------------------------------------------------------------------------------
# What goes inside a class block
# ---------------------------------------------------------------------------------------------
def block_text(s, by):
    kind: str = TYPE_SHORT.get(s["kind"], str(s["kind"]))
    if s["multi"] and s["kind"] == "Упражнение":
        kind += f" гр.{s['group']}"
    room = s["room"] or L["no_room"]
    flag = " ⚠" if any(w.startswith("requirement") for w in s["warnings"]) else ""
    if by == "cohort":
        second, third = short_list(s["teachers"], 2), f"{kind} · {room}{flag}"
    elif by == "teacher":
        second, third = short_list(s["cohorts"], 4), f"{kind} · {room}{flag}"
    else:
        second, third = short_list(s["cohorts"], 4), f"{kind} · {short_list(s['teachers'], 2)}{flag}"
    font = lambda **kw: InlineFont(rFont="Calibri", sz=FONT_SIZE, **kw)
    return CellRichText(TextBlock(font(b=True), s["subject"]),
                        TextBlock(font(), f"\n{second}\n"),
                        TextBlock(font(i=True), third))


def keys_of(s, by):
    if by == "cohort":
        return s["cohorts"]
    if by == "teacher":
        return s["teachers"]
    return [s["room"] or L["none_room"]]


def sort_key(name, by):
    if by == "cohort":
        m = re.match(r"^(.*?)\s*(\d+)$", name)
        return (int(m[2]), m[1]) if m else (99, name)
    return (0, name)


def sheet_name(name, used):
    base = re.sub(r"[\\/*?:\[\]]", "", name).strip()[:31] or "sheet"
    candidate, n = base, 2
    while candidate.lower() in used:
        candidate = f"{base[:28]}_{n}"
        n += 1
    used.add(candidate.lower())
    return candidate


# ---------------------------------------------------------------------------------------------
# Drawing one sheet
# ---------------------------------------------------------------------------------------------
def style_range(ws, row, c1, c2, fill=None, border=None, r2=None):
    for r in range(row, (r2 or row) + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(r, c)
            if fill is not None:
                cell.fill = fill
            if border is not None:
                cell.border = border


def merge(ws, r1, c1, r2, c2):
    if (r1, c1) != (r2, c2):
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)


def draw_sheet(ws, key, sessions, missing, by, grid):
    step, t0, ncols = grid["step"], grid["t0"], grid["ncols"]
    last_col = FIRST_COL + ncols - 1
    lunch0, lunch1 = LUNCH_START, LUNCH_START + LUNCH_MINUTES

    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 14
    for i in range(ncols):
        ws.column_dimensions[get_column_letter(FIRST_COL + i)].width = max(COL_WIDTH_HOUR * step / 60, 7)

    # --- title and legend
    ws.cell(1, 1, f"{L['title']} — {L['by'][by]}: {key}").font = Font(size=15, bold=True, color="1F3A5F")
    merge(ws, 1, 1, 1, last_col)
    ws.cell(2, 1, f"{L['period']}: {grid['first']:%d.%m.%Y} – {grid['last']:%d.%m.%Y}").font = Font(size=10, color="555555")
    merge(ws, 2, 1, 2, 4)
    span = 3 if step < 60 else 2
    col = 5
    for text, fill in ((L["lecture"], FILL["Лекция"]), (L["exercise"], FILL["Упражнение"]), (L["lunch"], LUNCH_FILL)):
        style_range(ws, 2, col, col + span - 1, fill=fill)
        ws.cell(2, col, text).alignment = Alignment(horizontal="center")
        ws.cell(2, col).font = Font(size=9)
        merge(ws, 2, col, 2, col + span - 1)
        col += span + 1
    ws.cell(2, col, L["warn"]).font = Font(size=9, color="B45309")
    ws.row_dimensions[1].height = 24

    # --- header row
    header = 4
    for c, text in ((1, L["week"]), (2, L["day"])):
        ws.cell(header, c, text)
    per = 60 // step if step < 60 and 60 % step == 0 else 1       # columns per header label
    for i in range(0, ncols, per):
        a = t0 + i * step
        ws.cell(header, FIRST_COL + i, f"{clock(a)}–{clock(a + per * step)}")
    for c in range(1, last_col + 1):
        cell = ws.cell(header, c)
        cell.fill, cell.font = HEAD_FILL, Font(bold=True, color="FFFFFF", size=9)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i in range(0, ncols, per):
        merge(ws, header, FIRST_COL + i, header, FIRST_COL + i + per - 1)
    ws.row_dimensions[header].height = 22
    ws.freeze_panes = ws.cell(header + 1, FIRST_COL)

    by_day = defaultdict(list)
    for s in sessions:
        by_day[s["day"]].append(s)

    # --- one row per lane per day
    row = header + 1
    for number, monday in enumerate(grid["weeks"], start=1):
        week_first = row
        for offset in range(7):
            day = monday + timedelta(days=offset)
            items = by_day.get(day, [])
            lane_of, lanes = assign_lanes(items)
            nrows = max(1, lanes)
            weekend = offset >= 5
            first_row, last_row = row, row + nrows - 1

            for r in range(first_row, last_row + 1):
                ws.row_dimensions[r].height = ROW_HEIGHT_CLASS if items else ROW_HEIGHT_EMPTY
                for i in range(ncols):
                    a = t0 + i * step
                    cell = ws.cell(r, FIRST_COL + i)
                    is_lunch = LUNCH_MINUTES > 0 and a < lunch1 and a + step > lunch0
                    cell.fill = LUNCH_FILL if is_lunch else (WEEKEND_FILL if weekend else PatternFill())
                    cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

            label = ws.cell(first_row, 2, f"{DAY_NAMES[offset]}\n{day:%d.%m.%Y}")
            label.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            label.font = Font(bold=True, size=9, color="B45309" if weekend else "1F3A5F")
            style_range(ws, first_row, 2, 2, fill=DAY_FILL, border=Border(left=THIN, right=THIN, top=THIN, bottom=THIN),
                        r2=last_row)
            merge(ws, first_row, 2, last_row, 2)

            for item in items:
                r = first_row + lane_of[id(item)]
                c1 = FIRST_COL + (item["start"] - t0) // step
                c2 = FIRST_COL + (item["end"] - t0) // step - 1
                edge = Side(style="thin", color=EDGE[item["kind"]])
                style_range(ws, r, c1, c2, fill=FILL[item["kind"]], border=Border(left=edge, right=edge, top=edge, bottom=edge))
                cell = ws.cell(r, c1)
                cell.value = block_text(item, by)
                cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
                if item["warnings"]:
                    note = Comment(f"{clock(item['start'])}–{clock(item['end'])}\n" + "\n".join(item["warnings"]), "timetable")
                    note.width, note.height = 300, 40 + 18 * len(item["warnings"])
                    cell.comment = note
                merge(ws, r, c1, r, c2)
            row = last_row + 1

        # week label and a heavy line between weeks
        wl = ws.cell(week_first, 1, number)
        wl.alignment = Alignment(horizontal="center", vertical="center")
        wl.font = Font(bold=True, size=12, color="1F3A5F")
        style_range(ws, week_first, 1, 1, fill=DAY_FILL, border=Border(left=THIN, right=THIN, top=THIN, bottom=THIN), r2=row - 1)
        merge(ws, week_first, 1, row - 1, 1)
        for c in range(1, last_col + 1):
            cell = ws.cell(row - 1, c)
            b = cell.border
            cell.border = Border(left=b.left, right=b.right, top=b.top, bottom=MEDIUM)

    # --- classes the solver could not place
    row += 1
    if missing:
        ws.cell(row, 1, f"{L['unplaced']} ({sum(m['hours'] for m in missing)} {L['hours'].lower()})")
        ws.cell(row, 1).font = Font(bold=True, size=12, color="B91C1C")
        merge(ws, row, 1, row, last_col)
        row += 1
        grouped = {}
        for m in missing:
            g = grouped.setdefault((m["subject"], m["kind"], m["group"] if m["multi"] else 0, tuple(m["teachers"])),
                                   dict(parts=0, hours=0, cohorts=m["cohorts"]))
            g["parts"] += 1
            g["hours"] += m["hours"]
        spans = [(L["subject"], 3, 7), (L["type"], 8, 8), (L["group"], 9, 9), (L["parts"], 10, 10),
                 (L["hours"], 11, 11), (L["teacher"] if by != "teacher" else L["cohorts"], 12, 15)]
        spans = [(t, a, min(b, last_col)) for t, a, b in spans if a <= last_col]
        for text, a, b in spans:
            ws.cell(row, a, text)
            style_range(ws, row, a, b, fill=HEAD_FILL)
            ws.cell(row, a).font = Font(bold=True, color="FFFFFF", size=9)
            merge(ws, row, a, row, b)
        row += 1
        for (subject, kind, group, teachers), g in sorted(grouped.items(), key=lambda kv: (-kv[1]["hours"], kv[0][0])):
            values = [subject, TYPE_SHORT.get(kind, kind), group or "",
                      g["parts"], g["hours"],
                      short_list(teachers, 3) if by != "teacher" else short_list(g["cohorts"], 4)]
            for (text, a, b), value in zip(spans, values):
                cell = ws.cell(row, a, value)
                cell.font = Font(size=9)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                style_range(ws, row, a, b, fill=FILL.get(kind) if a == 8 else None,
                            border=Border(left=THIN, right=THIN, top=THIN, bottom=THIN))
                merge(ws, row, a, row, b)
            ws.row_dimensions[row].height = 26
            row += 1

    # --- printing
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = f"{header}:{header}"
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.4


# ---------------------------------------------------------------------------------------------
def main(src="timetable.xlsx", dst="timetable_view.xlsx", by="cohort"):
    sessions, missing = read_sessions(src)
    if not sessions:
        raise SystemExit("The Timetable sheet is empty - nothing to draw.")
    grid = make_grid(sessions)

    placed_by, missing_by = defaultdict(list), defaultdict(list)
    for s in sessions:
        for key in keys_of(s, by):
            placed_by[key].append(s)
    if by != "room":
        for m in missing:
            for key in keys_of(dict(m, room=""), by):
                missing_by[key].append(m)
    keys = sorted(set(placed_by) | set(missing_by), key=lambda k: sort_key(k, by))

    wb = Workbook()
    overview = cast(Worksheet, wb.active)
    overview.title = L["overview"]
    used, names = {L["overview"].lower()}, {}
    for key in keys:
        names[key] = sheet_name(key, used)
        draw_sheet(wb.create_sheet(names[key]), key, placed_by.get(key, []), missing_by.get(key, []), by, grid)

    # --- overview sheet with links
    heads = [L["by"][by], L["placed_n"], L["placed_h"], L["unplaced_n"], L["unplaced_h"], L["placed_pct"]]
    for c, text in enumerate(heads, start=1):
        cell = overview.cell(1, c, text)
        cell.fill, cell.font = HEAD_FILL, Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    overview.row_dimensions[1].height = 32
    for r, key in enumerate(keys, start=2):
        ph = sum(s["hours"] for s in placed_by.get(key, []))
        mh = sum(m["hours"] for m in missing_by.get(key, []))
        link = cast(Cell, overview.cell(r, 1, key))
        link.hyperlink = f"#'{names[key]}'!A1"
        link.font = Font(color="0563C1", underline="single")
        overview.cell(r, 2, len(placed_by.get(key, [])))
        overview.cell(r, 3, ph)
        overview.cell(r, 4, len(missing_by.get(key, [])))
        overview.cell(r, 5, mh)
        pct = overview.cell(r, 6, ph / (ph + mh) if ph + mh else 1)
        pct.number_format = "0%"
        if mh:
            overview.cell(r, 5).font = Font(color="B91C1C", bold=True)
    overview.column_dimensions["A"].width = 28
    for c in "BCDEF":
        overview.column_dimensions[c].width = 16
    overview.freeze_panes = "A2"
    overview.page_setup.orientation = "portrait"
    overview.page_setup.fitToWidth, overview.page_setup.fitToHeight = 1, 0
    overview.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    overview.auto_filter.ref = f"A1:F{len(keys) + 1}"

    wb.save(dst)
    print(f"{len(keys)} sheets ({by}) written to {dst}")
    print(f"Grid: {clock(grid['t0'])}-{clock(grid['t1'])}, step {grid['step']} min, {len(grid['weeks'])} weeks")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("src", nargs="?", default="timetable.xlsx")
    parser.add_argument("dst", nargs="?", default="timetable_view.xlsx")
    parser.add_argument("--by", choices=["cohort", "teacher", "room"], default="cohort")
    args = parser.parse_args()
    main(args.src, args.dst, args.by)