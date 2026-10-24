"""Step 1 runner: import and validate the Excel file, print a summary, write import_report.xlsx.

Usage:  python run_import.py [input.xlsx] [report.xlsx]
"""
import math
import sys
from collections import Counter, defaultdict

import pandas as pd
from openpyxl.utils import get_column_letter

import config
import importer

LEVEL_ORDER = {"ERROR": 0, "WARNING": 1, "INFO": 2}


def capacity_hours():
    """Hours one teacher / one student cohort can possibly be in class during the period."""
    per_day = (config.DAY_END_MIN - config.DAY_START_MIN - config.LUNCH_MINUTES) // 60
    return config.NUM_DAYS * per_day


def load_table(sessions, resource_of, label):
    """Hours and sessions per teacher or cohort, with the share of the available time they need."""
    hours, count = Counter(), Counter()
    for s in sessions:
        for name in resource_of(s):
            hours[name] += s.length
            count[name] += 1
    cap = capacity_hours()
    df = pd.DataFrame({label: list(hours), "Hours": list(hours.values()), "Sessions": [count[n] for n in hours]})
    df["Load % of period"] = (100 * df["Hours"] / cap).round(0).astype(int)
    return df.sort_values("Hours", ascending=False).reset_index(drop=True)


def format_workbook(path):
    from openpyxl import load_workbook
    from openpyxl.styles import Font
    wb = load_workbook(path)
    for ws in wb:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 8), 70) # pyright: ignore[reportArgumentType]
    wb.save(path)


def main(path=config.INPUT_FILE, out="import_report.xlsx"):
    courses, issues, rows_read = importer.load_courses(path)
    sessions = importer.expand_sessions(courses)

    issue_df = pd.DataFrame([vars(i) for i in issues], columns=["level", "sheet", "row", "subject", "message"])
    issue_df = (issue_df.assign(_order=issue_df["level"].map(LEVEL_ORDER))
                .sort_values(["_order", "sheet", "row"]).drop(columns="_order"))

    course_df = pd.DataFrame([{
        "Type": c.kind, "Sheet": c.sheet, "Row": c.row, "Subject": c.subject,
        "Cohorts": ", ".join(c.cohorts), "Hours per group": c.hours, "Groups": c.groups,
        "Session length": c.duration, "Sessions": c.groups * math.ceil(c.hours / c.duration),
        "Teachers": ", ".join(c.teachers), "Rooms (best first)": ", ".join(c.rooms),
        "Requirements": "; ".join(r.raw for r in c.requirements),
    } for c in courses])

    teacher_df = load_table(sessions, lambda s: s.course.teachers, "Teacher")
    # Lower bound: only group 1 of each course, because the groups of one cohort may run in parallel.
    cohort_df = load_table([s for s in sessions if s.group == 1], lambda s: s.course.cohorts, "Cohort")

    room_first, room_any = Counter(), Counter()
    for c in courses:
        for i, room in enumerate(c.rooms):
            room_any[room] += 1
            room_first[room] += i == 0

    # Room load. Every session needs a room for its whole length, so a room can host at most `cap` hours.
    #   "Hours (only choice)" = hours of courses that list ONLY this room: nothing else can take them.
    #   "Hours (any choice)"  = hours of courses that list this room among their options.
    # Hours of parallel groups all count: each group sits in its own room.
    forced, listed = Counter(), Counter()
    set_hours = Counter()                     # room set -> hours of courses limited to exactly that set
    for s_ in sessions:
        rooms = s_.course.rooms
        if not rooms:
            continue
        set_hours[frozenset(rooms)] += s_.length
        for room in rooms:
            listed[room] += s_.length
        if len(rooms) == 1:
            forced[rooms[0]] += s_.length
    cap = capacity_hours()
    room_df = pd.DataFrame({"Room": list(room_any), "Courses listing it": list(room_any.values()),
                            "As first choice": [room_first[r] for r in room_any],
                            "Hours (only choice)": [forced[r] for r in room_any],
                            "Hours (any choice)": [listed[r] for r in room_any]})
    room_df["Only-choice load % of period"] = (100 * room_df["Hours (only choice)"] / cap).round(0).astype(int)
    room_df = room_df.sort_values("Only-choice load % of period", ascending=False).reset_index(drop=True)

    # Room groups: all courses whose room options lie inside a set of rooms must share those rooms.
    # If their hours exceed (number of rooms x capacity) the courses cannot all be placed, whatever the solver does.
    tight_rows = []
    for rset in set_hours:
        inside = sum(h for other, h in set_hours.items() if other <= rset)
        room_cap = len(rset) * cap
        tight_rows.append({"Rooms": ", ".join(sorted(rset)), "Rooms count": len(rset),
                           "Hours needing only these rooms": inside, "Capacity (hours)": room_cap,
                           "Load %": round(100 * inside / room_cap), "Over capacity by": max(0, inside - room_cap)})
    room_groups_df = (pd.DataFrame(tight_rows, columns=["Rooms", "Rooms count", "Hours needing only these rooms",
                                                        "Capacity (hours)", "Load %", "Over capacity by"])
                      .sort_values("Load %", ascending=False).reset_index(drop=True))

    with pd.ExcelWriter(out, engine="openpyxl") as xl:
        issue_df.to_excel(xl, sheet_name="Issues", index=False)
        course_df.to_excel(xl, sheet_name="Courses", index=False)
        teacher_df.to_excel(xl, sheet_name="Teachers", index=False)
        cohort_df.to_excel(xl, sheet_name="Cohorts", index=False)
        room_df.to_excel(xl, sheet_name="Rooms", index=False)
        room_groups_df.to_excel(xl, sheet_name="Room groups", index=False)
    format_workbook(out)

    levels = Counter(i.level for i in issues)
    cap = capacity_hours()
    print(f"Period: {config.START_DATE} .. {config.END_DATE} ({config.NUM_DAYS} days), "
          f"{config.DAY_START}-{config.DAY_END}, lunch {config.LUNCH_START} for {config.LUNCH_MINUTES} min "
          f"-> {cap} usable hours per teacher/cohort")
    print(f"Rows read: {dict(rows_read)}")
    print(f"Courses kept: {Counter(c.kind for c in courses)}  ->  {len(sessions)} sessions, "
          f"{sum(s.length for s in sessions)} hours to place")
    print(f"Issues: {dict(levels)}   (details in {out})")
    over_t = teacher_df[teacher_df["Hours"] > cap]
    over_c = cohort_df[cohort_df["Hours"] > cap]
    print(f"Teachers needing more than {cap} h (impossible as written): {len(over_t)}")
    print(f"Cohorts needing more than {cap} h even with parallel groups: {len(over_c)}")
    over_r = room_groups_df[room_groups_df["Over capacity by"] > 0]
    print(f"Room groups needing more than their rooms can hold (impossible as written): {len(over_r)}")
    print("Tightest room groups:\n" + room_groups_df.head(5).to_string(index=False))
    print("Busiest teachers:\n" + teacher_df.head(5).to_string(index=False))
    print("Busiest cohorts:\n" + cohort_df.head(5).to_string(index=False))


if __name__ == "__main__":
    main(*sys.argv[1:])